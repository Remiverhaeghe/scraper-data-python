"""Métiers qui recrutent en France, d'après l'enquête Besoins en Main-d'Œuvre (BMO)

Chaque année, France Travail interroge plus de 400 000 établissements sur leurs
projets de recrutement, par métier et par bassin d'emploi. L'enquête, publiée
en données ouvertes sur data.gouv.fr, est téléchargée à la demande puis résumée
par métier : France, région et département. Comparer deux années montre les
métiers en hausse ; la part des recrutements jugés difficiles, ceux où un
profil formé a le plus de chances.
"""

import io
import json
import os
import re

import openpyxl

from config import DATA_DIRECTORY
from job.fingerprint import normalize
from utils.api_client import request_bytes, request_json
from utils.logger import get_logger

logger = get_logger(__name__)

BMO_DATASET_URL = (
    "https://www.data.gouv.fr/api/1/datasets/enquete-besoins-en-main-doeuvre-bmo/"
)
BMO_DIRECTORY = DATA_DIRECTORY / "bmo"

# Le tableur d'une année pèse environ 4 Mo : au-delà, ce n'est plus l'enquête.
MAX_BMO_BYTES = 30 * 1024 * 1024
BMO_TIMEOUT = 120

# Titre des fichiers du jeu de données : « Enquête BMO 2026 »
BMO_TITLE_PATTERN = re.compile(r"BMO (\d{4})$")

# Colonnes lues ; met : projets de recrutement, xmet : jugés difficiles,
# smet : saisonniers. Une valeur masquée (« * ») compte pour zéro.
CODE_COLUMN = "Code métier BMO"
LABEL_COLUMN = "Nom métier BMO"
FAMILY_COLUMN = "Lbl_fam_met"
DEPARTMENT_COLUMN = "Dept"
DEPARTMENT_NAME_COLUMN = "NomDept"
COUNT_COLUMNS = ("met", "xmet", "smet")

# Seuils sous lesquels un taux n'a pas de sens : quelques projets seulement
MIN_FRANCE_PROJECTS = 500
MIN_DEPARTMENT_PROJECTS = 30
OUTLOOK_SIZE = 10
MINE_SIZE = 6

# Listes de la synthèse : clé, titre et explication
OUTLOOK_LISTS = (
    ("mine", "Vos métiers", "Rapprochés des intitulés de votre profil"),
    ("rising", "En hausse", "Projets de recrutement en plus depuis l'an dernier"),
    ("difficult", "Difficiles à pourvoir", "Employeurs en manque de candidats"),
    ("top", "Les plus demandés", "Le plus de projets de recrutement"),
)

# Mots trop vagues pour rapprocher un intitulé du profil d'un métier
PROFILE_STOP_WORDS = {
    "adjoint",
    "agent",
    "aujourd",
    "confirme",
    "experience",
    "junior",
    "polyvalent",
    "polyvalente",
    "profil",
    "responsable",
    "senior",
    "stage",
}
MIN_PROFILE_WORD = 5

# Racine commune minimale : « développeur » rejoint « développement »,
# « planificateur » rejoint « planification ».
MIN_COMMON_ROOT = 6


def read_count(pValue):
    """Retourne un nombre de projets, zéro pour une valeur masquée ou absente."""

    vText = str(pValue or "").strip()

    rCount = int(float(vText)) if re.fullmatch(r"\d+(\.0+)?", vText) else 0

    return rCount


def add_counts(pTotals, pCounts):
    """Ajoute des projets (total, difficiles, saisonniers) à un cumul."""

    rTotals = [vTotal + vCount for vTotal, vCount in zip(pTotals, pCounts)]

    return rTotals


def summarize_bmo(pContent, pYear):
    """Résume le tableur d'une année par métier : France et départements."""

    vBook = openpyxl.load_workbook(io.BytesIO(pContent), read_only=True)

    try:
        vSheet = max(vBook.worksheets, key=lambda pSheet: pSheet.max_row or 0)
        vRows = vSheet.iter_rows(values_only=True)
        vHeader = list(next(vRows))
        vIndex = {vName: vHeader.index(vName) for vName in vHeader if vName}
        vJobs = {}
        vDepartments = {}

        for vRow in vRows:
            vCode = vRow[vIndex[CODE_COLUMN]]

            if not vCode:
                continue

            vCounts = [read_count(vRow[vIndex[vColumn]]) for vColumn in COUNT_COLUMNS]
            vDepartment = str(vRow[vIndex[DEPARTMENT_COLUMN]] or "")
            vJob = vJobs.setdefault(
                vCode,
                {
                    "label": vRow[vIndex[LABEL_COLUMN]],
                    "family": vRow[vIndex[FAMILY_COLUMN]],
                    "france": [0, 0, 0],
                    "departments": {},
                },
            )
            vJob["france"] = add_counts(vJob["france"], vCounts)
            vJob["departments"][vDepartment] = add_counts(
                vJob["departments"].get(vDepartment, [0, 0, 0]), vCounts
            )
            vDepartments[vDepartment] = vRow[vIndex[DEPARTMENT_NAME_COLUMN]]

    finally:
        vBook.close()

    rSummary = {"year": pYear, "jobs": vJobs, "departments": vDepartments}

    return rSummary


def list_bmo_files():
    """Retourne l'adresse du tableur de chaque année publiée, par année."""

    vDataset = request_json("GET", BMO_DATASET_URL)

    rFiles = {
        int(vMatch.group(1)): vResource["url"]
        for vResource in vDataset.get("resources", [])
        if (vMatch := BMO_TITLE_PATTERN.search(str(vResource.get("title") or "")))
    }

    return rFiles


def get_summary_file(pYear):
    """Retourne le fichier du résumé d'une année."""

    rFile = BMO_DIRECTORY / f"bmo-{pYear}.json"

    return rFile


def refresh_labour_market(pYears=2):
    """Télécharge et résume les pYears dernières enquêtes absentes ; retourne les
    années disponibles. Lève requests.RequestException sans réseau.
    """

    vFiles = list_bmo_files()

    for vYear in sorted(vFiles, reverse=True)[:pYears]:
        if get_summary_file(vYear).exists():
            continue

        vContent = request_bytes(
            "GET",
            vFiles[vYear],
            pTimeout=BMO_TIMEOUT,
            pMaxSize=MAX_BMO_BYTES,
            pExpectedType="spreadsheetml",
        )
        vSummary = summarize_bmo(vContent, vYear)

        BMO_DIRECTORY.mkdir(parents=True, exist_ok=True)
        get_summary_file(vYear).write_text(
            json.dumps(vSummary, ensure_ascii=False), encoding="utf-8"
        )

        logger.info("Enquête BMO %s résumée : %s métiers", vYear, len(vSummary["jobs"]))

    rYears = list_summary_years()

    return rYears


def list_summary_years():
    """Retourne les années déjà résumées, la plus récente en tête."""

    rYears = sorted(
        (
            int(vFile.stem.removeprefix("bmo-"))
            for vFile in BMO_DIRECTORY.glob("bmo-*.json")
            if vFile.stem.removeprefix("bmo-").isdigit()
        ),
        reverse=True,
    )

    return rYears


def load_summary(pYear):
    """Retourne le résumé d'une année, ou None s'il n'a pas été téléchargé."""

    vFile = get_summary_file(pYear)

    rSummary = json.loads(vFile.read_text(encoding="utf-8")) if vFile.exists() else None

    return rSummary


def read_scope(pJob, pDepartment):
    """Retourne les projets d'un métier en France ou dans un département."""

    rCounts = (
        pJob["departments"].get(pDepartment, [0, 0, 0])
        if pDepartment
        else pJob["france"]
    )

    return rCounts


def find_profile_words(pProfile):
    """Retourne les mots significatifs des intitulés du profil, pour le rapprocher
    des métiers de l'enquête.
    """

    vTitles = [
        pProfile.get("headline") or "",
        *(vExperience.get("title") or "" for vExperience in pProfile["experiences"]),
        *(vVariant.get("name") or "" for vVariant in pProfile["cv_variants"]),
    ]

    rWords = {
        vWord
        for vTitle in vTitles
        for vWord in normalize(vTitle).split()
        if len(vWord) >= MIN_PROFILE_WORD and vWord not in PROFILE_STOP_WORDS
    }

    return rWords


def count_matching_words(pLabel, pWords):
    """Compte les mots du profil présents dans l'intitulé, par leur racine commune."""

    vLabelWords = normalize(pLabel).split()

    rCount = sum(
        any(
            len(os.path.commonprefix([vLabelWord, vWord]))
            >= min(MIN_COMMON_ROOT, len(vWord), len(vLabelWord))
            for vLabelWord in vLabelWords
            if len(vLabelWord) >= MIN_PROFILE_WORD
        )
        for vWord in pWords
    )

    return rCount


def build_outlook(pProfile, pDepartment=None):
    """Retourne les métiers à suivre dans le département donné, ou en France.

    Quatre listes : vos métiers, en hausse depuis l'an dernier, difficiles à
    pourvoir et les plus demandés. None tant qu'aucune enquête n'est téléchargée.
    """

    vYears = list_summary_years()

    if not vYears:
        return None

    vCurrent = load_summary(vYears[0])
    vPrevious = load_summary(vYears[1]) if len(vYears) > 1 else None
    vMinimum = MIN_DEPARTMENT_PROJECTS if pDepartment else MIN_FRANCE_PROJECTS
    vWords = find_profile_words(pProfile)
    vRows = []

    for vCode, vJob in vCurrent["jobs"].items():
        vProjects, vDifficult, vSeasonal = read_scope(vJob, pDepartment)

        if not vProjects:
            continue

        vBefore = (
            read_scope(vPrevious["jobs"][vCode], pDepartment)[0]
            if vPrevious and vCode in vPrevious["jobs"]
            else 0
        )

        vRows.append(
            {
                "code": vCode,
                "label": vJob["label"],
                "family": vJob["family"],
                "projects": vProjects,
                "difficult_share": round(100 * vDifficult / vProjects),
                "seasonal_share": round(100 * vSeasonal / vProjects),
                "change": (
                    round(100 * (vProjects - vBefore) / vBefore)
                    if vBefore >= vMinimum
                    else None
                ),
                "matching": count_matching_words(vJob["label"], vWords),
            }
        )

    vSolid = [vRow for vRow in vRows if vRow["projects"] >= vMinimum]

    rOutlook = {
        "year": vYears[0],
        "previous_year": vYears[1] if vPrevious else None,
        "scope": (
            vCurrent["departments"].get(pDepartment, pDepartment)
            if pDepartment
            else "France"
        ),
        "departments": sorted(vCurrent["departments"].items()),
        "mine": sorted(
            (vRow for vRow in vRows if vRow["matching"]),
            key=lambda pRow: (-pRow["matching"], -pRow["projects"]),
        )[:MINE_SIZE],
        "rising": sorted(
            (vRow for vRow in vSolid if vRow["change"] is not None),
            key=lambda pRow: -pRow["change"],
        )[:OUTLOOK_SIZE],
        "difficult": sorted(vSolid, key=lambda pRow: -pRow["difficult_share"])[
            :OUTLOOK_SIZE
        ],
        "top": sorted(vRows, key=lambda pRow: -pRow["projects"])[:OUTLOOK_SIZE],
    }

    return rOutlook


def describe_outlook_row(pRow):
    """Résume un métier : projets, part difficile et évolution sur un an."""

    vParts = [
        f"{pRow['projects']:,}".replace(",", " ") + " projets",
        f"{pRow['difficult_share']} % difficiles",
    ]

    if pRow["seasonal_share"] >= 50:
        vParts.append(f"{pRow['seasonal_share']} % saisonniers")

    if pRow["change"] is not None:
        vParts.append(f"{pRow['change']:+d} % sur un an")

    rText = " · ".join(vParts)

    return rText
